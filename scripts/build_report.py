#!/usr/bin/env python3
"""Fill the original course DOCX without changing its styles or page furniture.

The course DOCX is required as input and is never overwritten. Content is
inserted into its six existing chapters; original styles, section properties,
numbering, headers and footers are retained byte for byte. The cover logo uses
its original red image, with the template's green recoloring effects removed.
"""
import argparse
from copy import deepcopy
import csv
import hashlib
import json
from pathlib import Path
import re
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from lxml import etree

TITLE = '基于SSD的过滤向量检索系统课程实验'
REPO = 'https://github.com/Lydddddddd/fvs-ssd-filter-lab'
TEMPLATE_SHA256 = 'd6fdeebbe149e7063b48c3704b2a1e31f6e7e204eb53c49f0313354af4d3ec6b'
EDITABLE_PARTS = {
    '[Content_Types].xml', 'word/document.xml',
    'word/_rels/document.xml.rels', 'docProps/core.xml', 'docProps/app.xml',
}


def restore_cover_logo_color(document):
    """Remove the cover logo's recoloring while retaining its image and geometry."""
    for element in list(document._element.body)[:17]:
        for blip in element.iter(qn('a:blip')):
            if blip.get(qn('r:embed')) != 'rId9':
                continue
            for effect in list(blip):
                if effect.tag == qn('a:duotone'):
                    blip.remove(effect)
            for extension_list in blip.findall(qn('a:extLst')):
                for extension in list(extension_list):
                    if extension.find('{http://schemas.microsoft.com/office/drawing/2010/main}imgProps') is not None:
                        extension_list.remove(extension)


def replace_text(element, value):
    """Fill existing text nodes, retaining the paragraph, runs and controls."""
    nodes = list(element.iter(qn('w:t')))
    if not nodes:
        run = OxmlElement('w:r')
        text = OxmlElement('w:t'); run.append(text); element.append(run)
        nodes = [text]
    nodes[0].text = value
    for node in nodes[1:]:
        node.text = ''
    for node in list(element.iter(qn('w:showingPlcHdr'))):
        node.getparent().remove(node)
    for node in list(element.iter(qn('w:rStyle'))):
        if node.get(qn('w:val')) == 'aff2':
            node.getparent().remove(node)


def clone_text_paragraph(source, text):
    """Clone a source role rather than applying a new font or spacing preset."""
    result = deepcopy(source)
    for child in list(result):
        if child.tag != qn('w:pPr'):
            result.remove(child)
    for attr in list(result.attrib):
        if etree.QName(attr).localname in {'paraId', 'textId'}:
            del result.attrib[attr]
    run = OxmlElement('w:r')
    original_run = source.find(qn('w:r'))
    if original_run is not None and original_run.find(qn('w:rPr')) is not None:
        run.append(deepcopy(original_run.find(qn('w:rPr'))))
    parts = text.split('\n')
    for i, part in enumerate(parts):
        if i:
            run.append(OxmlElement('w:br'))
        node = OxmlElement('w:t'); node.text = part
        node.set(qn('xml:space'), 'preserve'); run.append(node)
    result.append(run)
    return result


def fill_bound_properties(payload, name, identity):
    tree = etree.fromstring(payload)
    values = {'title': TITLE, 'creator': identity['name'],
              'category': identity['department'], 'Company': identity['class'],
              'keywords': identity['student_id'], 'lastModifiedBy': identity['name']}
    for node in tree.iter():
        key = etree.QName(node).localname
        if key in values:
            node.text = values[key]
    return etree.tostring(tree, encoding='UTF-8', xml_declaration=True,
                          standalone=True)


def import_refreshed_toc(output, refreshed):
    """Import Word's TOC cache and targets without importing its style rewrites."""
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    def style_names(payload):
        tree = etree.fromstring(payload)
        return {s.get(qn('w:styleId')): s.find(qn('w:name')).get(qn('w:val')).casefold()
                for s in tree.findall(qn('w:style')) if s.find(qn('w:name')) is not None}
    def text(element):
        return ''.join(element.xpath('.//w:t/text()', namespaces=ns))
    with zipfile.ZipFile(output) as original, zipfile.ZipFile(refreshed) as word:
        parts = {n: original.read(n) for n in original.namelist()}
        tree = etree.fromstring(parts['word/document.xml'])
        word_tree = etree.fromstring(word.read('word/document.xml'))
        original_names = style_names(parts['word/styles.xml'])
        word_names = style_names(word.read('word/styles.xml'))
        ids_by_name = {name: sid for sid, name in original_names.items()}
        content = deepcopy(word_tree.find('.//' + qn('w:body')).find(qn('w:sdt')).find(qn('w:sdtContent')))
        for tag in ['pStyle', 'rStyle', 'tblStyle']:
            for node in content.iter(qn('w:'+tag)):
                sid = node.get(qn('w:val'))
                node.set(qn('w:val'), ids_by_name[word_names[sid]])
        target = tree.find(qn('w:body')).find(qn('w:sdt'))
        target.replace(target.find(qn('w:sdtContent')), content)
        paragraphs = {text(p): p for p in tree.find(qn('w:body')).findall(qn('w:p'))}
        existing = {n.get(qn('w:name')) for n in tree.iter(qn('w:bookmarkStart'))}
        next_id = max(int(n.get(qn('w:id'))) for n in tree.iter(qn('w:bookmarkStart'))) + 1
        targets = {n.get(qn('w:anchor')) for n in content.iter(qn('w:hyperlink'))}
        for paragraph in word_tree.find(qn('w:body')).findall(qn('w:p')):
            for mark in paragraph.findall(qn('w:bookmarkStart')):
                name = mark.get(qn('w:name'))
                if name not in targets or name in existing:
                    continue
                destination = paragraphs[text(paragraph)]
                start = OxmlElement('w:bookmarkStart'); end = OxmlElement('w:bookmarkEnd')
                start.set(qn('w:id'), str(next_id)); start.set(qn('w:name'), name)
                end.set(qn('w:id'), str(next_id)); next_id += 1
                destination.insert(1 if destination.find(qn('w:pPr')) is not None else 0, start)
                destination.append(end); existing.add(name)
        assert targets <= existing
        parts['word/document.xml'] = etree.tostring(tree, encoding='UTF-8', xml_declaration=True, standalone=True)
    with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED) as final:
        for name, payload in parts.items():
            final.writestr(name, payload)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--template', type=Path, required=True)
    ap.add_argument('--project', type=Path, required=True)
    ap.add_argument('--output', type=Path, required=True)
    ap.add_argument('--identity', type=Path, required=True,
                    help='Local JSON with name, student_id, class, department')
    ap.add_argument('--refreshed-docx', type=Path,
                    help='A temporary Word copy with its TOC updated; only the field cache is imported')
    args = ap.parse_args(); root = args.project
    if args.template.resolve() == args.output.resolve():
        raise ValueError('Keep the course template unchanged; use a distinct output path')
    template_sha = hashlib.sha256(args.template.read_bytes()).hexdigest()
    if template_sha != TEMPLATE_SHA256:
        raise ValueError('This report fills the fixed course template; verify its version')
    identity = json.loads(args.identity.read_text())
    data = json.loads((root/'figures/analysis.json').read_text())
    rows = list(csv.DictReader((root/'data/baseline/results.csv').open()))
    bs = {(r['system'], r['bucket'], r['L']): r for r in data['baseline_summary']}
    os = {(r['variant'], r['bucket'], r['L']): r for r in data['optimization_summary']}
    doc = Document(args.template)
    restore_cover_logo_color(doc)
    original = list(doc._element.body)
    body_source = deepcopy(original[22])
    heading_source = deepcopy(original[23])
    anchors = {1: original[21], 2: original[26], 3: original[35],
               4: original[40], 5: original[42], 6: original[45]}
    secondary = {(1, 1): original[23], (2, 1): original[27],
                 (2, 2): original[29], (2, 3): original[31]}
    # Remove only empty body placeholders. Preserve the original section
    # paragraphs, chapter headings, bookmarks and chapter page breaks.
    for i in [24, 25, 28, 30, 32, 33, 43, 46]:
        original[i].getparent().remove(original[i])
    for row, value in zip(doc.tables[1].rows,
                          [identity['department'], identity['class'],
                           identity['name'], identity['student_id']]):
        replace_text(row.cells[1]._tc, value)
    cursor = original[21]; chapter = 0; table_count = 0
    used_intro = False
    headings = []

    def place(element):
        nonlocal cursor
        cursor.addnext(element); cursor = element
        return element

    def p(text='', indent=True, keep=False):
        nonlocal used_intro, cursor
        if chapter == 1 and not used_intro:
            obj = Paragraph(original[22], doc._body)
            replace_text(obj._p, text); cursor = obj._p; used_intro = True
        else:
            obj = Paragraph(place(clone_text_paragraph(body_source, text)), doc._body)
        if not indent:
            pr = obj._p.get_or_add_pPr(); ind = pr.find(qn('w:ind'))
            if ind is None:
                ind = OxmlElement('w:ind'); pr.append(ind)
            ind.set(qn('w:firstLine'), '0'); ind.set(qn('w:firstLineChars'), '0')
        if keep:
            obj.paragraph_format.keep_with_next = True
        if REPO in text:
            # Keep typography inherited from the source paragraph; add a real link.
            before, _, after = text.partition(REPO)
            for run in list(obj._p.findall(qn('w:r'))):
                obj._p.remove(run)
            obj.add_run(before)
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), doc.part.relate_to(REPO, RT.HYPERLINK, is_external=True))
            run = OxmlElement('w:r'); node = OxmlElement('w:t'); node.text = REPO
            run.append(node); link.append(run); obj._p.append(link); obj.add_run(after)
        return obj

    def h(text, level=1, newpage=False):
        nonlocal cursor, chapter
        if level == 1:
            chapter = int(text.split()[0]); cursor = anchors[chapter]
            obj = Paragraph(cursor, doc._body)
            # The original six chapter titles remain unchanged.
        else:
            section_number = int(text.split()[0].split('.')[1])
            anchor = secondary.get((chapter, section_number))
            if anchor is not None:
                obj = Paragraph(anchor, doc._body); cursor = anchor
            else:
                title = text.partition(' ')[2] if chapter <= 2 else text
                obj = Paragraph(place(clone_text_paragraph(heading_source, title)), doc._body)
                if chapter >= 3:
                    pr = obj._p.get_or_add_pPr(); num = OxmlElement('w:numPr')
                    node = OxmlElement('w:numId'); node.set(qn('w:val'), '0')
                    num.append(node); pr.append(num)
        headings.append({'chapter': chapter, 'level': level, 'title': text})
        return obj

    def caption(text, keep=False):
        obj = p(text, False)
        pr = obj._p.get_or_add_pPr()
        for child in list(pr):
            pr.remove(child)
        st = OxmlElement('w:pStyle'); st.set(qn('w:val'), doc.styles['Caption'].style_id)
        pr.append(st)
        if keep:
            obj.paragraph_format.keep_with_next = True
        return obj

    def table(title, headers, values, widths=None):
        nonlocal table_count
        table_count += 1
        caption(re.sub(r'^表\d+', '表'+str(table_count), title), True)
        t = doc.add_table(rows=1, cols=len(headers)); place(t._tbl)
        t.style = doc.styles['Table Grid']; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        if widths is None:
            widths = [14.4/len(headers)]*len(headers)
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
        for cell, txt in zip(t.rows[0].cells, headers):
            cell.text = str(txt)
        for values_row in values:
            for cell, txt in zip(t.add_row().cells, values_row):
                cell.text = str(txt)
        for i, row in enumerate(t.rows):
            trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if i == 0:
                trpr.append(OxmlElement('w:tblHeader'))
            for j, cell in enumerate(row.cells):
                cell.width = Cm(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.paragraph_format.first_line_indent = 0
                    pr = para._p.get_or_add_pPr(); ind = pr.find(qn('w:ind'))
                    ind.set(qn('w:firstLineChars'), '0')
                    para.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    if i == 0:
                        para.paragraph_format.keep_with_next = True
                        for run in para.runs:
                            run.bold = True
        return t

    def figure(name, label):
        obj = p('', False, True); obj.alignment = WD_ALIGN_PARAGRAPH.CENTER
        obj.add_run().add_picture(str(root/'figures'/name), width=Cm(14.4))
        caption(label)

    def code(lines):
        obj = p('\n'.join(lines), False)
        obj.alignment = WD_ALIGN_PARAGRAPH.LEFT
        obj.paragraph_format.keep_together = True
        return obj

    def f(row, metric, digits=2):
        return f"{row[metric+'_median']:.{digits}f}"

    h('1 实验环境与复现结果')
    p('本实验在统一的 YFCC 多标签向量负载上，比较 PipeANN-Filter 与 GateANN 的召回率、查询时延、逻辑 I/O 和峰值内存，并实现 GateANN 的访问位图优化。测试显示，PipeANN-Filter 在强过滤场景中通过切换执行路径获得较高召回率；GateANN 在 High 负载中的查询时延较低。访问位图进一步将 GateANN 在 Low 桶中的查询时延中位数降低14.18%，Recall@10保持34.30%。')
    h('1.1 实验环境',2)
    table('表1 实验平台与固定软件版本',['项目','实际配置'],[
      ['CPU','2 × AMD EPYC 7542；64个物理核、128个逻辑CPU'],
      ['内存','宿主机可见约251 GiB'],
      ['SSD与文件系统','Samsung 990 EVO Plus 4TB，NVMe；/workspace位于nvme0n1p1，ext4'],
      ['操作系统','Ubuntu 22.04.4 LTS，x86-64；内核5.15.0-139-generic'],
      ['编译环境','GCC/G++ 11.4.0，CMake Release，Linux AIO'],
      ['PipeANN提交','2ec84fe20c37dc3cc236b66f7eba31c0cc80bf1e'],
      ['GateANN提交','71e87cbf0e25c0400c0109dd25a64a7c4dccc9dd'],
      ['基础实验编号','20261029-083358；每组3次独立进程运行'],
      ['公共查询参数','threads=1，beamwidth=8，K=10，mem_L=0，cache_budget=0'],
      ['GateANN结构','FullAdj邻居上限32；不使用GPU'],
    ],[3.0,11.4])
    h('1.2 复现结果',2)
    p('通过 findmnt 确认，实验工作目录位于990 EVO Plus的ext4分区。使用课程脚本编译固定版本源码，并为GateANN应用AIO兼容补丁。两套Debug索引均成功构建；High正确性验证的Recall@10分别为96.20%和95.85%，检查器通过。基础实验编号为20261029-083358，优化代码在独立的GateANN工作树中实现，原始结果、运行命令和补丁保存在项目仓库中。')

    h('1.3 调试记录',2)
    table('表9 关键问题及处理记录',['现象','定位与处理','处理结果'],[
      ['SSH断开、下载中断','修复Tailscale服务资源异常并调整SSH保活；利用下载脚本断点续传','两套正式索引下载完成并通过校验'],
      ['GateANN编译兼容','课程AIO兼容补丁将control block数组访问修正为逐项访问','固定commit、Release+AIO，Debug正确性检查通过'],
      ['补丁初次无法应用','上游YFCC驱动使用CRLF；保留原文件换行后重新生成统一差分','完整补丁反向apply检查通过；源码diff检查通过'],
      ['top-10比对有细小差异','记录实际返回数量并核验标签与L2排序；重复同一hash程序10次','语义检查通过；固定程序重复运行出现0.05个百分点Recall波动'],
      ['High个别轮次耗时较高','汇总7轮中位数、最小值和最大值','图3展示逐轮波动，Low和Medium的收益较稳定'],
    ],[3.0,6.2,5.2])
    p('基础实验18次、补充实验42次查询调用均正常退出。命令、日志和CSV按实验分别保存，AIO兼容补丁与访问位图补丁均可复现。')

    h('2 系统设计理解')
    h('2.1 PipeANN-Filter',2)
    p('PipeANN 将图搜索的计算与SSD访问流水化，在候选处理、距离计算和异步读取之间安排重叠，减少串行等待。PipeANN-Filter在这一框架中加入过滤：DRAM保存PQ近似表示和概率过滤摘要，SSD保存完整向量、图及属性索引。搜索利用摘要选出候选超集，再验证真实属性，以少量额外探索减少属性读取。')
    p('查询入口spec_filter_search.cpp根据选择率和过滤精度估计代价，在Pre、In、Post三种路径中选择执行方式。Pre先利用标签缩小候选集合，In在图遍历过程中判断过滤条件，Post先扩大搜索再筛选。概率摘要允许候选阶段出现假阳性，最终的真实标签验证保证返回结果满足过滤条件。')
    h('2.2 GateANN',2)
    p('GateANN 将图遍历与完整向量获取分开：DRAM中的标签CSR用于检查过滤条件，PQ用于估计距离，FullAdj提供节点邻接表；完整向量及磁盘图位于SSD。Pre-I/O Filtering在发起SSD请求前判断标签。对于不匹配的节点，Graph Tunneling利用内存邻接表和PQ距离继续扩展，穿过这些节点到达满足条件的区域，保持过滤后的图可达性。')
    p('这一设计省去了许多无效向量的 SSD 读取，但保留了邻居去重、PQ 距离计算和候选池更新。课程 Formal FullAdj 文件为 132000024 字节，约 125.89 MiB，此外还需标签和 PQ 结构。减少 I/O 后，遍历中的 CPU 与内存访问可能成为主要开销；内存邻接表也会随向量规模和邻居数量增长。这正是本实验选择访问集合为优化对象的背景。')
    h('2.3 设计比较',2)
    table('表2 两种设计的主要差异',['系统','过滤位置与可达性','主要资源与潜在代价'],[
      ['PipeANN-Filter','成本模型选择Pre/In/Post；概率摘要允许候选超集，最终验证属性','DRAM摘要与PQ；SSD向量、图、属性索引；可能增加误判探索和磁盘索引空间'],
      ['GateANN','发起I/O前检查；FullAdj和PQ支持穿过不匹配节点','DRAM标签CSR、PQ、FullAdj；省向量读取，仍付遍历与内存代价'],
    ],[3.0,5.7,5.7])
    p('PipeANN-Filter以路由和摘要控制候选；GateANN在I/O前过滤，并借FullAdj保留图可达性。前者偏重SSD索引，后者付出DRAM代价，比较时需同时看Recall与时延。')

    h('3 实验方法')
    h('3.1 数据与过滤查询',2)
    p('使用课程dataset-v1.0的YFCC真实多标签数据。向量为192维uint8，距离为L2；查询要求指定标签属于数据点的标签集合。Debug使用10万向量进行构建与排错，Formal使用100万向量及课程预构建索引，各自对应相同规模的精确Ground Truth。Recall@10表示返回结果命中过滤后精确top-10的比例，再对200次查询求均值。')
    table('表3 Formal测试负载及选择理由',['桶','选择率范围','选择率中位数','过滤强度与查询数'],[
      ['Low','0.250%～0.851%','0.545%','强过滤；200'],
      ['Medium','1.765%～4.995%','2.958%','中等过滤；200'],
      ['High','7.074%～17.089%','12.323%','弱过滤；200'],
    ],[2.0,4.2,3.2,5.0])
    p('每桶包含5个真实标签，每个标签40次查询。选择这三档负载，是为了观察有效候选从少到多时，系统的过滤路径、无效遍历和SSD读取如何变化。桶间同时包含标签频率、语义和向量分布差异，刻画了实际多标签查询的性能变化。')
    h('3.2 参数控制与统计方法',2)
    p('基础实验固定AIO、单查询线程、beamwidth=8、K=10、mem_L=0、缓存预算0，GateANN FullAdj邻居上限32。High分别运行PipeANN的L=20/40/80和GateANN的L=160/320/640；Medium、Low采用代表性L=40和320。两系统内部预算尺度不同，比较时以Recall相近的配置对齐检索质量，再考察查询时延和资源消耗。')
    p('每系统每轮5个配置，两系统合计10条；独立重复3轮，共30行。按system、bucket、L分组，以中位数表示典型性能，图中误差线表示三轮最小值至最大值。系统专用指标分别用于解释其内部行为。')
    p('时延使用查询阶段均值：PipeANN取程序报告值，GateANN按适配器以1000000/QPS换算。逻辑I/O采用程序内部的读取计数；峰值RSS覆盖查询子进程的加载、初始化和查询阶段，以KiB记录、MiB展示。GateANN的IO、Tunnel和Process计时分别对应轮询处理、内存穿行与节点处理代码，用于定位开销。')
    p('实验在共享服务器上运行，操作系统页面缓存采用默认状态。补充测试固定CPU 127并交替执行两种模式，开始与结束的一分钟load average分别约22.32和23.27。基础比较与优化对照分开汇总，优化收益由同一时段的hash与bitmap对照计算。')
    h('3.3 主要执行命令',2)
    code(['bash scripts/prepare_pipeann.sh','bash scripts/prepare_gateann.sh',
          'bash scripts/run_debug_validation.sh','bash scripts/download_indexes.sh',
          'bash scripts/check_indexes.sh','bash scripts/run_formal_experiments.sh'])
    p('Debug阶段分别调用两个适配器的build子命令，设置tier=debug、构建线程8。正式运行脚本自动执行3轮，查询线程固定为1。两套正式索引均通过SHA-256及manifest校验；基础实验的18个查询调用全部正常退出，汇总CSV包含完整的30条结果。')

    h('4 实验结果与机制分析')
    h('4.1 High负载下的召回率与时延',2)
    figure('01_recall_latency.png','图1 High负载的Recall与时延（3轮中位数及范围）')
    values=[]
    for s,ls in [('pipeann',['20','40','80']),('gateann',['160','320','640'])]:
        for L in ls:
            r=bs[(s,'high',L)]
            values.append(['PipeANN-Filter' if s=='pipeann' else 'GateANN',L,f(r,'recall_percent'),f"{r['mean_latency_us_median']/1000:.3f}",f(r,'qps',1),f(r,'mean_ios')])
    table('表4 High负载的性能中位数',['系统','L','Recall %','时延 ms','QPS','逻辑I/O'],values,[3.5,1.3,2.4,2.4,2.5,2.3])
    p('增加搜索预算后，两系统的Recall和逻辑I/O均上升。PipeANN的逻辑I/O从94.60增至134.00、222.40；GateANN从24.64增至46.74、88.89。在约81%的Recall水平，GateANN L=160达到81.40%和0.624 ms，PipeANN L=20达到80.90%和2.291 ms，前者的时延中位数约为后者的1/3.67。')
    p('在约93%的Recall水平，GateANN L=320为93.25%、1.125 ms，PipeANN L=40为92.20%、1.257 ms，时延差距缩小。更高预算下，GateANN L=640达到98.65%、2.119 ms，PipeANN L=80达到97.30%、2.302 ms。GateANN在本次High负载中的时延优势，随检索质量提高而收窄。')
    p('PipeANN L=20的三轮时延为1.077～2.331 ms，中位数高于L=40；L=40的范围为1.256～1.346 ms。与I/O随预算增长的稳定趋势相比，查询时延还受到运行状态的影响。图1同时展示中位数和范围，呈现了这组非单调变化及其轮次波动。')
    h('4.2 过滤强度与执行路径',2)
    figure('02_selectivity_performance.png','图2 代表性预算下的选择率与性能（3轮中位数及范围）')
    values=[]
    for s,L in [('pipeann','40'),('gateann','320')]:
        for b in ['low','medium','high']:
            r=bs[(s,b,L)]
            values.append(['PipeANN-Filter' if s=='pipeann' else 'GateANN',b.title(),f(r,'recall_percent'),f"{r['mean_latency_us_median']/1000:.3f}",f(r,'mean_ios')])
    table('表5 代表性预算下的性能中位数',['系统','桶','Recall %','时延 ms','逻辑I/O'],values,[3.4,2.3,2.9,2.9,2.9])
    table('表6 查询路径和内部行为的中位数',['桶','PipeANN Pre/In/Post','Gate跳过数','Gate IO/Tunnel/Process μs'],[
      [b.title(),'/'.join(f(bs[('pipeann',b,'40')],m,0) for m in ['pre_filter_queries','in_filter_queries','post_filter_queries']),
       f(bs[('gateann',b,'320')],'filter_skips'),
       '/'.join(f(bs[('gateann',b,'320')],m,0) for m in ['io_time_us','tunnel_time_us','process_time_us'])]
      for b in ['low','medium','high']],[2.0,4.0,2.7,5.7])
    p('PipeANN在Low的200个查询全部走Pre路径，在Medium和High全部走In路径。Low达到98.60% Recall和0.542 ms，Medium为91.20%和2.676 ms。路由计数与性能变化对应：Pre先缩小标签候选集合，适合Low中较少的有效候选；In路径需要在图遍历中持续处理候选和过滤判断。')
    p('GateANN在Low中每查询平均读取4.38次，时延为0.747 ms，Tunnel计时中位数为529 μs，Recall为34.30%。High L=320则为46.74次读取、1.125 ms和93.25% Recall。Low中的有效向量读取很少，内存中对大量不匹配节点的扩展成为值得优化的开销。')
    p('为解释Low的低Recall，补充记录了pipe_search的实际返回数量。200次查询共返回757个ID，平均每次3.785个；25次凑齐10个，22次无结果。其余槽位由驱动填0。按K=10计算，757个真实返回结果对应37.85%的Recall上限，与实测34.30%接近。固定L=320时，有效答案数量不足是Low低召回的重要原因，后续可围绕搜索预算和终止条件改进。')
    h('4.3 空间与内存代价',2)
    resource=[]
    for s,label,extra in [('pipeann','PipeANN-Filter','过滤属性索引及较大的磁盘图'),('gateann','GateANN','标签CSR、PQ、约125.89 MiB FullAdj')]:
        rr=[r for r in rows if r['system']==s]
        resource.append([label,f"{float(rr[0]['index_bytes'])/1024**3:.3f}",f"{min(float(r['peak_rss_kb']) for r in rr)/1024:.2f}～{max(float(r['peak_rss_kb']) for r in rr)/1024:.2f}",extra])
    table('表7 原始正式实验的资源代价',['系统','索引 GiB','进程峰值RSS MiB','主要结构'],resource,[3.0,2.0,4.0,5.4])
    p('按适配器统计的索引文件，PipeANN占8.425 GiB，GateANN占0.475 GiB。对应的查询进程峰值RSS分别约147～149 MiB和411～415 MiB。GateANN以更多DRAM保存标签、PQ和邻接信息，换取紧凑的磁盘索引及较少的向量读取。RSS还包含加载阶段的临时数组和查询缓冲，体现了进程整体的内存需求。')
    h('4.4 关键现象与瓶颈判断',2)
    p('第一个关键现象是PipeANN在Low中更准、更快：98.60%/0.542 ms，对应Medium的91.20%/2.676 ms。两桶分别采用Pre和In路径，说明有效候选规模和执行路线共同影响检索成本。Medium中较重的图探索和过滤处理，是后续分析In路径的切入点。')
    p('第二个关键现象是GateANN在Low中I/O少而内存穿行耗时较多。每查询4.38次读取、529 μs Tunnel和3.785个真实结果，分别揭示了读取节省、遍历成本与答案不足。本实验选择内存遍历中的访问集合为优化对象。')

    h('5 优化方向设计')
    h('5.1 观察与机制',2)
    p('Low中每查询平均跳过320.69个不匹配节点，Tunnel计时中位数为529 μs。沿src/search/pipe_search.cpp追踪可见，普通邻居扩展和Graph Tunneling都会先对visited执行find，再对首次出现的节点执行insert。这类重复哈希访问发生在遍历热路径上，而节点ID连续分布在[0,N)内，适合改用直接寻址的位图。')
    h('5.2 优化方案与代码实现',2)
    p('增加查询内的精确位图，每个节点对应1 bit。访问时计算64位字编号和位掩码，检查并置位，保持“首次发现才加入候选”的语义。GATEANN_VISITED_MODE=bitmap启用位图，hash保留原路径。每次查询独立分配并清零位图，访问标记由当前查询单独维护。')
    code(['auto &word = words_[id >> 6];',
          'const uint64_t mask = uint64_t{1} << (id & 63);',
          'const bool fresh = (word & mask) == 0;',
          'word |= mask;',
          'return fresh;'])
    p('实现位于optimization/dense_visited.h，完整上游修改位于patches/gateann_dense_visited.patch，覆盖普通扩展、tunneling和初始节点标记，并增加可选的实际返回数量导出。非法ID会抛出异常，避免越界访问。查询驱动的外层墙钟计时包含位图分配、初始化和释放；性能运行关闭结果导出。所有代码与复现入口：'+REPO+'。')
    h('5.3 验证方案与正确性',2)
    p('以同一个Release+AIO构建中的hash和bitmap两条路径对照，改变访问集合实现，保持CPU 127、数据、索引、搜索预算及其他查询参数一致。每桶运行7轮，High每轮包含3档L；桶顺序按种子20261029打乱，两种模式按轮次交替AB/BA，合计产生70条结果。')
    p('使用AddressSanitizer和UndefinedBehaviorSanitizer检查集合实现，完成1200030次与std::unordered_set的差分测试，覆盖字边界、不同节点规模、查询隔离及非法ID。再对Debug/Formal共10组配置，按真实返回数量核验标签条件、ID唯一性、精确L2排序，并重算Recall，全部通过。')
    p('9组配置的有序top-10文件逐字节一致，另一组Recall相差0.05个百分点。对同一hash程序的Debug High L=160再运行10次，也观察到85.75%和85.80%两种Recall。五个正式配置中，hash与bitmap的Recall中位数均相同。')
    h('5.4 性能结果与预期检验',2)
    figure('03_optimization.png','图3 同构建hash与bitmap对照（7轮中位数及范围）')
    values=[]
    for b,L in [('high','160'),('high','320'),('high','640'),('medium','320'),('low','320')]:
        a,z=os[('hash',b,L)],os[('bitmap',b,L)]
        imp=next(x for x in data['improvements'] if (x['bucket'],x['L'])==(b,L))
        values.append([b.title()+'/'+L,f"{a['mean_latency_us_median']:.2f}",f"{z['mean_latency_us_median']:.2f}",f"{imp['latency_reduction_percent']:.2f}%",f"{z['recall_percent_median']:.2f}"])
    table('表8 优化实验的查询时延中位数',['配置','hash μs','bitmap μs','时延降幅','两组Recall %'],values,[3.1,2.8,2.8,2.6,3.1])
    p('优化预期是减少内存遍历中的集合访问成本，尤其适合无效节点多、SSD读取少的强过滤负载。Low的时延中位数从851.79 μs降至730.99 μs，降低14.18%，加速比1.165；Tunnel从606 μs降至533 μs，降低12.05%，逻辑I/O为4.38，Recall为34.30%。Medium时延降低13.01%。这两个桶的7轮配对中，bitmap均更快，两组时延范围也分别错开。')
    p('High L=160/320/640的时延中位数分别降低9.45%、9.95%、10.49%。其中L=320和640各有一轮bitmap更慢，七轮平均时延的最大值分别为3.484 ms和6.452 ms，图3的误差线展示了这两次波动。与High相比，Low和Medium的逐轮收益更稳定。')
    h('5.5 代价与适用范围',2)
    p('位图占ceil(N/64)×8字节。100万节点时，每查询占125000字节，即122.07 KiB，并承担O(N/64)的清零成本。Low中hash与bitmap的进程峰值RSS中位数分别约402.24和402.26 MiB。位图分配、清零及释放均计入查询驱动的外层计时，收益来自用这部分连续存储替代遍历中的哈希访问。')
    p('该方案适合ID稠密、节点规模适中的负载。节点数增大时，清零成本随之增加；并发查询的位图空间也按并发数增长。更大规模可采用分段位图或代际标记。针对Low的有效答案不足问题，还可结合实际返回数量设计自适应搜索预算，同时观察新增I/O、时延和Recall。')

    h('6 实验总结与结论')
    p('PipeANN-Filter通过成本模型选择过滤路径，Low采用Pre后达到98.60% Recall和0.542 ms时延。GateANN通过预I/O过滤与内存tunneling减少无效向量读取，在High的相近Recall水平取得较低时延。两者分别体现了过滤路径选择与内存信息组织对SSD检索的作用。')
    p('针对GateANN遍历中的重复哈希访问，本实验实现了查询内位图，并用7轮对照验证效果。五组配置的时延中位数降低9.45%～14.18%，Recall中位数相同；Low和Medium在全部配对轮次中均更快。优化以每查询122.07 KiB的位图空间减少集合访问成本，适合节点ID稠密的检索负载。')
    p('本次测试说明，SSD过滤检索需要同时考虑有效答案数量、I/O访问和内存遍历。Low中平均返回3.785个真实答案，提示下一步可将访问位图与自适应预算结合，继续改善强过滤场景的检索质量与效率。')

    args.output.parent.mkdir(parents=True, exist_ok=True)
    temp = args.output.with_suffix('.working.docx'); doc.save(temp)
    changed = []
    with zipfile.ZipFile(args.template) as source, zipfile.ZipFile(temp) as created:
        names = list(source.namelist())
        names.extend(n for n in created.namelist() if n not in names)
        with zipfile.ZipFile(args.output, 'w', compression=zipfile.ZIP_DEFLATED) as final:
            for name in names:
                if name in {'docProps/core.xml','docProps/app.xml'}:
                    payload = fill_bound_properties(source.read(name), name, identity)
                elif name in source.namelist() and name not in EDITABLE_PARTS:
                    payload = source.read(name)
                else:
                    payload = created.read(name)
                if name in source.namelist() and payload != source.read(name):
                    changed.append(name)
                final.writestr(name, payload)
    temp.unlink()
    if args.refreshed_docx:
        import_refreshed_toc(args.output, args.refreshed_docx)
    with zipfile.ZipFile(args.template) as source, zipfile.ZipFile(args.output) as final:
        preserved = [name for name in source.namelist() if source.read(name) == final.read(name)]
        for name in set(source.namelist()) - EDITABLE_PARTS:
            assert source.read(name) == final.read(name), name
        before = Document(args.template); after = Document(args.output)
        assert len(before.sections) == len(after.sections) == 4
        for a, b in zip(before.sections, after.sections):
            assert etree.tostring(a._sectPr) == etree.tostring(b._sectPr)
    audit = {'template_sha256': template_sha,
             'docx_sha256': hashlib.sha256(args.output.read_bytes()).hexdigest(),
             'section_count': 4, 'source_styles_unchanged': True,
             'body_font': '宋体 / Times New Roman (original Normal style)',
             'body_font_pt': 12, 'line_spacing_multiple': 1.5,
             'cover_logo_color': 'original red; template recoloring effects removed',
             'original_parts_preserved': preserved, 'changed_parts': changed,
             'headings': headings, 'tables': table_count, 'figures': 3,
             'source_rows': {'baseline': len(rows), 'optimization': 70},
             'bibliography_included': False,
             'toc': ('Refreshed in Word; original TOC control, styles and target bookmarks retained'
                     if args.refreshed_docx else 'Original TOC field retained; refresh in Word before PDF export')}
    args.output.with_suffix('.audit.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2))
    print(json.dumps({k:v for k,v in audit.items() if k not in {'headings','original_parts_preserved'}}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
